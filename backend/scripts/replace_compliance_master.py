"""Convert the SME workbook and optionally replace the local MORAX master.

The source workbook describes most legal triggers in prose. Monthly rules use
the approved MORAX default (previous-month period, due on day 15 of the new
month); annual and event-driven rules remain MANUAL until an SME supplies an
exact machine-safe statutory date.

Run from the backend directory:
    python scripts/replace_compliance_master.py
    python scripts/replace_compliance_master.py --apply
"""

from __future__ import annotations

import argparse
import json
import re
import sqlite3
import sys
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import create_engine, delete, select
from sqlalchemy.orm import Session


BACKEND_DIR = Path(__file__).resolve().parents[1]
REPO_DIR = BACKEND_DIR.parent
sys.path.insert(0, str(BACKEND_DIR))


def initial_effective_date() -> date:
    """Make a newly loaded active master eligible for the closed month.

    Monthly obligations generated in the current month represent the previous
    month. Starting an imported active rule on the import day would therefore
    suppress the first valid occurrence. The source workbook has no separate
    commencement date, so the conservative operational default is the first
    day of the immediately preceding month.
    """
    previous_month_end = date.today().replace(day=1) - timedelta(days=1)
    return previous_month_end.replace(day=1)

from app.models import (  # noqa: E402
    ApplicabilityEvaluation,
    AuditLog,
    ComplianceAssignment,
    ComplianceComment,
    ComplianceEvidence,
    ComplianceImport,
    ComplianceImportError,
    ComplianceImportRow,
    ComplianceInstance,
    ComplianceRule,
    ComplianceRuleVersion,
    ComplianceSubmission,
    ComplianceWorkflowHistory,
    InAppNotification,
    IndustryType,
    Organization,
    RuleApplicabilityCriteria,
    State,
    User,
)
from app.schemas import RuleInput  # noqa: E402


HEADERS = [
    "Compliance ID",
    "Compliance Name",
    "Description",
    "Act",
    "Rule",
    "Section",
    "Compliance Type",
    "Document Type",
    "Form Number",
    "Legal Description",
    "Consequence / Penalty",
    "Version",
    "Entity Type",
    "State",
    "Industry Type",
    "Employee Threshold",
    "Unit Compliance Types",
    "Frequency",
    "Applicable Months",
    "Due Date Rule",
    "Due Date Offset",
    "Due Date Anchor",
    "Required Document",
    "Risk Level",
    "Effective From",
    "Effective To",
    "Active",
]

DOCUMENT_TYPES = {
    "procedure": "PROCEDURAL",
    "registers": "REGISTER",
    "records": "RECORDS",
    "remittance": "REMITTANCE",
    "returns": "RETURN",
    "return": "RETURN",
    "display": "DISPLAY",
}


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def normalized_frequency(value: Any) -> str:
    frequency = clean(value).lower()
    if "half-year" in frequency:
        return "HALF_YEARLY"
    if frequency.startswith("monthly"):
        return "MONTHLY"
    if frequency == "annual" or frequency.startswith("annual /"):
        return "ANNUAL"
    if frequency.startswith("one-time"):
        return "ONE_TIME"
    return "EVENT_BASED"


def unit_compliance_types(sheet_name: str, applicability: Any, name: str, rule_reference: str) -> list[str]:
    """Translate the SME wording into the Unit's controlled applicability tags."""
    text = f"{name} {rule_reference}".upper()
    specialized: list[str] = []
    if "EPF" in text or "PROVIDENT FUND" in text:
        specialized.append("PF")
    if "ESI" in text or "ESIC" in text:
        specialized.append("ESI")
    if specialized:
        return specialized
    # Rows in the dedicated Shops worksheet belong to that Unit profile unless
    # the rule was explicitly identified as PF/ESI above. Do this before text
    # classification so ordinary words such as "determine" cannot be mistaken
    # for a Mines trigger.
    if sheet_name == "SHOPS":
        return ["Shops & Establishments"]
    value = clean(applicability).upper()
    if "BOCW" in value:
        return ["BOCW"]
    if re.search(r"\bMINES?\b", value) or re.search(r"\bMINES?\b", text):
        return ["Mines"]
    # ALL COVERED and other Central establishment rules are universal within
    # their entity/threshold criteria rather than tied to a UI checkbox.
    return []


def is_tamil_nadu_act(act: str) -> bool:
    return "TAMIL NADU" in act.upper()


def targets_for(sheet_name: str, applicability: Any) -> list[tuple[str, str]]:
    if sheet_name == "SHOPS":
        return [("UNIT", "Commercial")]
    value = clean(applicability).upper()
    if value == "CONTRACT LABOUR":
        return [("CONTRACTOR", "")]
    if value == "BOCW":
        return [("CONTRACTOR_SITE", "Construction")]
    if "BOCW" in value and "INDUSTRIAL ESTABLISHMENTS" in value:
        # Applicability is OR-based in the source; one normalized MORAX rule is
        # required per subject/industry because a criterion row is conjunctive.
        return [("UNIT", "Factory"), ("CONTRACTOR_SITE", "Construction")]
    return [("UNIT", "Factory")]


def source_rows(source: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    for sheet_name in ("Central", "SHOPS"):
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Required source worksheet is missing: {sheet_name}")
        sheet = workbook[sheet_name]
        headers = [clean(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        current_act = ""
        sequence = 0
        for values in sheet.iter_rows(min_row=2, values_only=True):
            row = {headers[index]: values[index] for index in range(min(len(headers), len(values)))}
            serial = values[0] if values else None
            name = clean(values[1] if len(values) > 1 else None)
            if serial in (None, "") and name:
                current_act = name
                continue
            if not isinstance(serial, (int, float)) or not name:
                continue
            sequence += 1
            if sheet_name == "Central":
                applicability = row.get("APPLICABITY")
                rule_reference = clean(row.get("Form No. / Rule Reference"))
                evidence = clean(row.get("Evidence / Records to Verify"))
                document_type = clean(row.get("Document Type"))
                trigger = clean(row.get("Due Date / Trigger"))
                source_frequency = clean(row.get("Frequency"))
                risk = clean(row.get("Risk Level")) or "MEDIUM"
                threshold = row.get("Limit")
                status_logic = clean(row.get("Logic"))
            else:
                applicability = "SHOPS & ESTABLISHMENTS"
                rule_reference = clean(row.get("Form No. / Rule Reference"))
                evidence = clean(row.get("Evidence / Records to Verify"))
                document_type = clean(row.get("Document Type"))
                trigger = clean(row.get("Due Date / Trigger"))
                source_frequency = clean(row.get("Frequency"))
                risk = clean(row.get("Risk Level")) or "MEDIUM"
                threshold = None
                status_logic = clean(row.get("Due Date Logic (Software Rule)"))

            targets = targets_for(sheet_name, applicability)
            for target_index, (entity_type, industry) in enumerate(targets, start=1):
                suffix = "" if len(targets) == 1 else f"-{target_index}"
                compliance_id = f"TN-{sheet_name}-{sequence:03d}{suffix}"
                legal_parts = [
                    f"Source frequency: {source_frequency}" if source_frequency else "",
                    f"Due date / trigger: {trigger}" if trigger else "",
                    f"Source status logic: {status_logic}" if status_logic else "",
                ]
                records.append(
                    {
                        "Compliance ID": compliance_id,
                        "Compliance Name": name[:250],
                        "Description": f"Evidence / records to verify: {evidence}" if evidence else "",
                        "Act": current_act[:250],
                        "Rule": rule_reference[:250],
                        "Section": "",
                        "Compliance Type": clean(applicability)[:100],
                        "Document Type": DOCUMENT_TYPES.get(document_type.lower(), "RECORDS"),
                        "Form Number": "",
                        "Legal Description": "\n".join(part for part in legal_parts if part),
                        "Consequence / Penalty": "",
                        "Version": 1,
                        "Entity Type": entity_type,
                        "State": "Tamil Nadu" if is_tamil_nadu_act(current_act) else "",
                        "Industry Type": "",
                        "Employee Threshold": int(threshold) if threshold not in (None, "") else "",
                        "Unit Compliance Types": ",".join(unit_compliance_types(sheet_name, applicability, name, rule_reference)) if entity_type == "UNIT" else "",
                        "Frequency": normalized_frequency(source_frequency),
                        "Applicable Months": "",
                        "Due Date Rule": "FIXED_DAY_OF_MONTH" if normalized_frequency(source_frequency) == "MONTHLY" else "MANUAL",
                        "Due Date Offset": 15 if normalized_frequency(source_frequency) == "MONTHLY" else "",
                        "Due Date Anchor": "",
                        "Required Document": evidence[:250],
                        "Risk Level": risk.upper(),
                        "Effective From": initial_effective_date().isoformat(),
                        "Effective To": "",
                        "Active": "TRUE",
                    }
                )
    if not records:
        raise ValueError("The source workbook contains no compliance rows")
    identities = [item["Compliance ID"] for item in records]
    if len(identities) != len(set(identities)):
        raise ValueError("The converted workbook contains duplicate Compliance IDs")
    return records


def write_product_workbook(records: list[dict[str, Any]], output: Path, source: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Compliance Master"
    sheet.append(HEADERS)
    for record in records:
        sheet.append([record[header] for header in HEADERS])
    header_fill = PatternFill("solid", fgColor="34452A")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 34
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(max(max(len(clean(cell.value)) for cell in column) + 2, 12), 42)
        sheet.column_dimensions[letter].width = width
    notes = workbook.create_sheet("Import Notes")
    notes.append(["Item", "Value"])
    notes.append(["Source workbook", source.name])
    notes.append(["Converted rows", len(records)])
    notes.append(["State mapping", "Central Acts = All states; Tamil Nadu Acts = Tamil Nadu"])
    notes.append(["Applicability", "Unit Compliance Types are the primary filter; industry remains unrestricted unless the SME source specifies it"])
    notes.append(["Due-date handling", "Monthly = previous month, due day 15; narrative annual/event rules remain MANUAL"])
    notes.append(["Threshold", "Central Limit is preserved as Employee Threshold"])
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 90
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def validate_records(records: list[dict[str, Any]], states: dict[str, str], industries: dict[str, str]) -> list[RuleInput]:
    payloads: list[RuleInput] = []
    for row in records:
        state_name = clean(row["State"])
        state_id = states.get(state_name.lower()) if state_name else None
        industry_name = clean(row["Industry Type"])
        industry_id = industries.get(industry_name.lower()) if industry_name else None
        if state_name and not state_id:
            raise ValueError(f"State is not configured: {row['State']}")
        if industry_name and not industry_id:
            raise ValueError(f"Industry Type is not configured: {industry_name}")
        payloads.append(
            RuleInput(
                compliance_id=row["Compliance ID"],
                name=row["Compliance Name"],
                description=row["Description"] or None,
                act=row["Act"] or None,
                rule_reference=row["Rule"] or None,
                section=None,
                compliance_type=row["Compliance Type"] or None,
                document_type=row["Document Type"] or None,
                form_number=None,
                legal_description=row["Legal Description"] or None,
                consequence_or_penalty=None,
                version=1,
                entity_type=row["Entity Type"],
                state_id=state_id,
                industry_type_id=industry_id,
                employee_threshold=row["Employee Threshold"] or None,
                unit_compliance_types=[item.strip() for item in str(row["Unit Compliance Types"] or "").split(",") if item.strip()],
                frequency=row["Frequency"],
                applicable_months=[],
                due_date_rule=row["Due Date Rule"],
                due_date_offset=row["Due Date Offset"] or None,
                required_document=row["Required Document"] or None,
                risk_level=row["Risk Level"],
                effective_from=date.fromisoformat(row["Effective From"]),
                active=True,
            )
        )
    return payloads


def database_catalog(database: Path) -> tuple[dict[str, str], dict[str, str]]:
    connection = sqlite3.connect(database)
    try:
        states = {name.lower(): identifier for identifier, name in connection.execute("SELECT id, name FROM states")}
        industries = {name.lower(): identifier for identifier, name in connection.execute("SELECT id, name FROM industry_types")}
        return states, industries
    finally:
        connection.close()


def backup_database(database: Path) -> Path:
    backup_dir = BACKEND_DIR / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = backup_dir / f"morax-before-compliance-master-{timestamp}.db"
    source_connection = sqlite3.connect(database)
    target_connection = sqlite3.connect(backup)
    try:
        source_connection.backup(target_connection)
    finally:
        target_connection.close()
        source_connection.close()
    return backup


def replace_master(database: Path, source: Path, records: list[dict[str, Any]], payloads: list[RuleInput]) -> tuple[Path, dict[str, int]]:
    backup = backup_database(database)
    engine = create_engine(f"sqlite:///{database.as_posix()}", connect_args={"check_same_thread": False})
    counts: dict[str, int] = {}
    with Session(engine) as db:
        try:
            master_org = db.scalar(select(Organization).where(Organization.is_system.is_(True)).order_by(Organization.created_at, Organization.id))
            actor = db.scalar(select(User).where(User.platform_role == "MORAX_ADMIN").order_by(User.created_at, User.id))
            if not master_org or not actor:
                raise RuntimeError("MORAX platform master organization or administrator is not configured")
            # The local MORAX-DEMO workspace normally restores visual QA data
            # at startup. Once an SME master is explicitly installed, suppress
            # that automatic seed so QA rules cannot contaminate the master.
            master_org.development_reset_at = datetime.now(UTC)

            old_rule_ids = list(db.scalars(select(ComplianceRule.id).where(ComplianceRule.organization_id == master_org.id)))
            old_version_ids = list(db.scalars(select(ComplianceRuleVersion.id).where(ComplianceRuleVersion.rule_id.in_(old_rule_ids)))) if old_rule_ids else []
            instance_ids = list(db.scalars(select(ComplianceInstance.id).where(ComplianceInstance.rule_version_id.in_(old_version_ids)))) if old_version_ids else []
            counts["rules_removed"] = len(old_rule_ids)
            counts["versions_removed"] = len(old_version_ids)
            counts["instances_removed"] = len(instance_ids)

            if instance_ids:
                db.execute(delete(InAppNotification).where(InAppNotification.reference_id.in_(instance_ids)))
                for model in (ComplianceEvidence, ComplianceComment, ComplianceSubmission, ComplianceAssignment, ComplianceWorkflowHistory):
                    db.execute(delete(model).where(model.instance_id.in_(instance_ids)))
                db.execute(delete(ComplianceInstance).where(ComplianceInstance.id.in_(instance_ids)))
            if old_version_ids:
                db.execute(delete(ApplicabilityEvaluation).where(ApplicabilityEvaluation.rule_version_id.in_(old_version_ids)))
                db.execute(delete(RuleApplicabilityCriteria).where(RuleApplicabilityCriteria.rule_version_id.in_(old_version_ids)))
                db.execute(delete(ComplianceRuleVersion).where(ComplianceRuleVersion.id.in_(old_version_ids)))
            if old_rule_ids:
                db.execute(delete(ComplianceRule).where(ComplianceRule.id.in_(old_rule_ids)))

            old_import_ids = list(db.scalars(select(ComplianceImport.id).where(ComplianceImport.organization_id == master_org.id)))
            if old_import_ids:
                import_row_ids = list(db.scalars(select(ComplianceImportRow.id).where(ComplianceImportRow.import_id.in_(old_import_ids))))
                if import_row_ids:
                    db.execute(delete(ComplianceImportError).where(ComplianceImportError.import_row_id.in_(import_row_ids)))
                db.execute(delete(ComplianceImportRow).where(ComplianceImportRow.import_id.in_(old_import_ids)))
                db.execute(delete(ComplianceImport).where(ComplianceImport.id.in_(old_import_ids)))

            import_record = ComplianceImport(
                organization_id=master_org.id,
                uploaded_by_id=actor.id,
                original_filename=source.name,
                status="CONFIRMED",
                total_rows=len(records),
                valid_rows=len(records),
                error_rows=0,
                confirmed_at=datetime.now(UTC),
            )
            db.add(import_record)
            db.flush()

            for row_number, (record, payload) in enumerate(zip(records, payloads, strict=True), start=2):
                db.add(ComplianceImportRow(import_id=import_record.id, row_number=row_number, payload_json=json.dumps(record), valid=True))
                rule = ComplianceRule(organization_id=master_org.id, compliance_id=payload.compliance_id, created_by_id=actor.id)
                db.add(rule)
                db.flush()
                version = ComplianceRuleVersion(
                    rule_id=rule.id,
                    version=payload.version,
                    name=payload.name,
                    description=payload.description,
                    act=payload.act,
                    rule_reference=payload.rule_reference,
                    section=payload.section,
                    compliance_type=payload.compliance_type,
                    document_type=payload.document_type,
                    form_number=payload.form_number,
                    legal_description=payload.legal_description,
                    consequence_or_penalty=payload.consequence_or_penalty,
                    frequency=payload.frequency.upper(),
                    applicable_months_json="[]",
                    due_date_rule=payload.due_date_rule.upper(),
                    due_date_offset=payload.due_date_offset,
                    due_date_anchor=None,
                    grace_days=0,
                    required_document=payload.required_document,
                    risk_level=payload.risk_level.upper(),
                    effective_from=payload.effective_from,
                    effective_to=None,
                    active=True,
                )
                db.add(version)
                db.flush()
                db.add(
                    RuleApplicabilityCriteria(
                        rule_version_id=version.id,
                        entity_type=payload.entity_type,
                        state_id=payload.state_id,
                        industry_type_id=payload.industry_type_id,
                        employee_threshold=payload.employee_threshold,
                        unit_compliance_types_json=json.dumps(payload.unit_compliance_types),
                    )
                )

            db.add(
                AuditLog(
                    organization_id=master_org.id,
                    actor_id=actor.id,
                    action="REPLACE_COMPLIANCE_MASTER",
                    module="COMPLIANCE_MASTER",
                    entity_type="ComplianceImport",
                    entity_id=import_record.id,
                    new_value_json=json.dumps({"source": source.name, "imported": len(records), **counts}),
                    reason="Replace the development compliance master with the supplied SME workbook",
                )
            )
            db.commit()
            counts["rules_imported"] = len(records)
        except Exception:
            db.rollback()
            raise
    engine.dispose()
    return backup, counts


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=REPO_DIR / "docs" / "complieance_master_due_date_status_logic.xlsx")
    parser.add_argument("--output", type=Path, default=REPO_DIR / "docs" / "compliance_master_product_template.xlsx")
    parser.add_argument("--database", type=Path, default=BACKEND_DIR / "morax.db")
    parser.add_argument("--apply", action="store_true", help="Replace the local database master after creating a backup")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    source = args.source.resolve()
    output = args.output.resolve()
    database = args.database.resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    if not database.is_file():
        raise FileNotFoundError(database)
    records = source_rows(source)
    write_product_workbook(records, output, source)
    states, industries = database_catalog(database)
    payloads = validate_records(records, states, industries)
    source_count = sum(1 for record in records if not str(record["Compliance ID"]).endswith("-2"))
    print(json.dumps({"mode": "apply" if args.apply else "dry-run", "source_rows": source_count, "normalized_rows": len(records), "output": str(output)}, indent=2))
    if args.apply:
        backup, counts = replace_master(database, source, records, payloads)
        print(json.dumps({"backup": str(backup), **counts}, indent=2))


if __name__ == "__main__":
    main()
