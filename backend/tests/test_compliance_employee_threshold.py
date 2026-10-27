import asyncio
import io
import json
from datetime import date
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.api.v1 import router as api
from app.db.base import Base
from app.models import ComplianceImport, ComplianceImportRow, ComplianceRule, ComplianceRuleVersion, IndustryType, Organization, RuleApplicabilityCriteria, State, User
from app.schemas import ComplianceActDeleteInput, ComplianceActStatusInput, ComplianceMasterDeleteInput
from app.services.compliance import evaluate


def test_employee_threshold_is_part_of_rule_applicability():
    version = SimpleNamespace(id="rule-version")
    criteria = SimpleNamespace(
        entity_type="UNIT",
        state_id="TN",
        industry_type_id="FACTORY",
        employee_threshold=10,
        applicability_flag=True,
        notes=None,
    )

    below_threshold = SimpleNamespace(
        entity_type="UNIT",
        state_id="TN",
        industry_type_id="FACTORY",
        total_employees=9,
    )
    at_threshold = SimpleNamespace(
        entity_type="UNIT",
        state_id="TN",
        industry_type_id="FACTORY",
        total_employees=10,
    )

    applicable, explanation = evaluate(version, criteria, below_threshold)
    assert applicable is False
    assert explanation["matched"]["employee_threshold"] is False

    applicable, explanation = evaluate(version, criteria, at_threshold)
    assert applicable is True
    assert explanation["matched"]["employee_threshold"] is True


def test_compliance_master_can_be_filtered_and_listed_by_act():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        organization = Organization(name="MORAX Master", code="MORAX-MASTER", status="ACTIVE", is_system=True)
        admin = User(name="Product Admin", email="product-admin@example.test", password_hash="hash", active=True, platform_role="MORAX_ADMIN")
        db.add_all([organization, admin])
        db.flush()
        for sequence, act in enumerate(("Code on Wages", "Industrial Relations Code"), start=1):
            rule = ComplianceRule(organization_id=organization.id, compliance_id=f"ACT-{sequence:03d}", created_by_id=admin.id)
            db.add(rule)
            db.flush()
            version = ComplianceRuleVersion(
                rule_id=rule.id,
                version=1,
                name=f"Rule {sequence}",
                act=act,
                frequency="EVENT_BASED",
                due_date_rule="MANUAL",
                risk_level="MEDIUM",
                effective_from=date(2026, 1, 1),
                active=True,
            )
            db.add(version)
            db.flush()
            db.add(RuleApplicabilityCriteria(rule_version_id=version.id, entity_type="UNIT"))
        db.commit()

        assert api.list_compliance_acts(user=admin, db=db) == ["Code on Wages", "Industrial Relations Code"]
        result = api.list_rules(
            paginated=True,
            page=1,
            page_size=50,
            q=None,
            act="Code on Wages",
            frequency=None,
            active=None,
            sort_by="act",
            sort_dir="asc",
            user=admin,
            db=db,
        )
        assert result["total"] == 1
        assert result["items"][0]["act"] == "Code on Wages"


def test_compliance_master_export_replace_and_confirmed_deletion():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        organization = Organization(name="MORAX Master", code="MORAX-MASTER", status="ACTIVE", is_system=True)
        admin = User(name="Product Admin", email="master-admin@example.test", password_hash="hash", active=True, platform_role="MORAX_ADMIN")
        state = State(name="Tamil Nadu", code="TN")
        industry = IndustryType(name="Factory", code="FACTORY")
        db.add_all([organization, admin, state, industry])
        db.flush()

        old_rule = ComplianceRule(organization_id=organization.id, compliance_id="OLD-001", created_by_id=admin.id)
        db.add(old_rule)
        db.flush()
        old_version = ComplianceRuleVersion(rule_id=old_rule.id, version=1, name="Old rule", act="Old Act", frequency="EVENT_BASED", due_date_rule="MANUAL", risk_level="LOW", effective_from=date(2025, 1, 1), active=True)
        db.add(old_version)
        db.flush()
        db.add(RuleApplicabilityCriteria(rule_version_id=old_version.id, entity_type="UNIT"))

        record = ComplianceImport(organization_id=organization.id, uploaded_by_id=admin.id, original_filename="replacement.xlsx", status="VALIDATED", total_rows=1, valid_rows=1, error_rows=0)
        db.add(record)
        db.flush()
        payload = {
            "Compliance ID": "NEW-001", "Compliance Name": "New rule", "Description": "Evidence guidance",
            "Act": "New Act", "Rule": "Rule 1", "Section": "", "Compliance Type": "STATUTORY",
            "Document Type": "RECORDS", "Form Number": "", "Legal Description": "Manual trigger",
            "Consequence / Penalty": "", "Version": 1, "Entity Type": "UNIT", "State": "Tamil Nadu",
            "Industry Type": "Factory", "Employee Threshold": 10, "Frequency": "EVENT_BASED",
            "Applicable Months": "", "Due Date Rule": "MANUAL", "Due Date Offset": "",
            "Due Date Anchor": "", "Required Document": "Evidence", "Risk Level": "HIGH",
            "Effective From": "2026-01-01", "Effective To": "", "Active": "TRUE",
        }
        db.add(ComplianceImportRow(import_id=record.id, row_number=2, payload_json=json.dumps(payload), valid=True))
        db.commit()

        result = api.confirm_import(record.id, mode="REPLACE", confirmation="REPLACE", user=admin, db=db)
        assert result["mode"] == "REPLACE"
        assert result["imported"] == 1
        assert db.query(ComplianceRule).count() == 1
        new_rule = db.query(ComplianceRule).one()
        assert new_rule.compliance_id == "NEW-001"

        response = api.export_compliance_master(user=admin, db=db)

        async def response_bytes():
            chunks = []
            async for chunk in response.body_iterator:
                chunks.append(chunk.encode() if isinstance(chunk, str) else chunk)
            return b"".join(chunks)

        workbook = load_workbook(io.BytesIO(asyncio.run(response_bytes())), read_only=True, data_only=True)
        exported = workbook["Compliance Master"]
        assert exported.max_row == 2
        assert exported["A2"].value == "NEW-001"
        assert exported["D2"].value == "New Act"

        status = api.set_compliance_act_status(ComplianceActStatusInput(act="New Act", active=False), user=admin, db=db)
        assert status["updated"] == 1
        assert db.get(ComplianceRuleVersion, db.query(ComplianceRuleVersion.id).scalar()).active is False

        with pytest.raises(HTTPException, match="exact Compliance ID"):
            api.delete_compliance_rule(new_rule.id, ComplianceMasterDeleteInput(confirmation="wrong"), user=admin, db=db)
        deleted = api.delete_compliance_rule(new_rule.id, ComplianceMasterDeleteInput(confirmation="NEW-001"), user=admin, db=db)
        assert deleted["deleted"]["rules"] == 1
        assert db.query(ComplianceRule).count() == 0

        for sequence in range(2):
            rule = ComplianceRule(organization_id=organization.id, compliance_id=f"BULK-{sequence}", created_by_id=admin.id)
            db.add(rule); db.flush()
            version = ComplianceRuleVersion(rule_id=rule.id, version=1, name=f"Bulk {sequence}", act="Bulk Act", frequency="EVENT_BASED", due_date_rule="MANUAL", risk_level="LOW", effective_from=date(2026, 1, 1), active=True)
            db.add(version); db.flush()
            db.add(RuleApplicabilityCriteria(rule_version_id=version.id, entity_type="UNIT"))
        db.commit()
        with pytest.raises(HTTPException, match="exact Act name"):
            api.delete_compliance_act(ComplianceActDeleteInput(act="Bulk Act", confirmation="wrong"), user=admin, db=db)
        deleted_act = api.delete_compliance_act(ComplianceActDeleteInput(act="Bulk Act", confirmation="Bulk Act"), user=admin, db=db)
        assert deleted_act["deleted"]["rules"] == 2
        assert db.query(ComplianceRule).count() == 0
